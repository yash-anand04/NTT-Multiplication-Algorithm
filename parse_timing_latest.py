import openpyxl
from pathlib import Path

for name in [
    r"Reports/Table.xlsx",
    r"Reports/Table_new.xlsx",
    r"Reports/Table_new2.xlsx",
    r"Reports/Table_new3.xlsx",
]:
    wb = openpyxl.load_workbook(Path(name), data_only=True)
    ws = wb.active
    print(name)
    print("  max_row", ws.max_row, "max_col", ws.max_column)
    print("  WNS", ws['B2'].value)
    for r in range(2, min(7, ws.max_row + 1)):
        print("   ", r, ws[f'B{r}'].value, ws[f'F{r}'].value, ws[f'G{r}'].value)
    wb.close()
